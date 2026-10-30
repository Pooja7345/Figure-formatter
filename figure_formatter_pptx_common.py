
import os
import tempfile
import uuid

import cv2
from PIL import Image
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.enum.text import MSO_ANCHOR, MSO_AUTO_SIZE
from pptx.dml.color import RGBColor


FONT_NAME = "Times New Roman"
FONT_SIZE_PT = 9

MAX_WIDTH_IN = 6.0
MAX_HEIGHT_IN = 9.0

VERTICAL_ASPECT_THRESHOLD = 2.2


def is_vertical_region(region):
    x0, y0, x1, y1 = region["box"]

    w = max(x1 - x0, 1)
    h = max(y1 - y0, 1)

    return h > VERTICAL_ASPECT_THRESHOLD * w


def new_presentation_with_background(
    cleaned_image_np,
    slide_width_in=MAX_WIDTH_IN,
    max_height_in=MAX_HEIGHT_IN
):
    """
    Create a PowerPoint presentation with the cleaned image
    as a full-slide background.

    Preserves the original image aspect ratio within
    maximum dimensions of 6 x 9 inches.
    """

    img_h, img_w = cleaned_image_np.shape[:2]

    px_to_in = min(
        slide_width_in / img_w,
        max_height_in / img_h
    )

    final_width_in = img_w * px_to_in
    final_height_in = img_h * px_to_in

    prs = Presentation()

    prs.slide_width = Inches(final_width_in)
    prs.slide_height = Inches(final_height_in)

    slide = prs.slides.add_slide(prs.slide_layouts[6])

    tmp_fd, tmp_path = tempfile.mkstemp(suffix=".png")
    os.close(tmp_fd)

    try:
        Image.fromarray(
            cv2.cvtColor(cleaned_image_np, cv2.COLOR_BGR2RGB)
        ).save(tmp_path)

        slide.shapes.add_picture(
            tmp_path,
            Inches(0),
            Inches(0),
            width=Inches(final_width_in),
            height=Inches(final_height_in)
        )

    finally:
        if os.path.exists(tmp_path):
            os.remove(tmp_path)

    return prs, slide, px_to_in


def add_label_textbox(slide, region, text, px_to_in):
    """
    Add editable Times New Roman text.
    Automatically handles tall/narrow vertical text regions.
    """

    x0, y0, x1, y1 = region["box"]

    vertical = is_vertical_region(region)

    box_w_px = x1 - x0
    box_h_px = y1 - y0

    cx_in = ((x0 + x1) / 2) * px_to_in
    cy_in = ((y0 + y1) / 2) * px_to_in

    if vertical:
        unrot_w_in = box_h_px * px_to_in
        unrot_h_in = box_w_px * px_to_in
    else:
        unrot_w_in = box_w_px * px_to_in
        unrot_h_in = box_h_px * px_to_in

    left = Inches(cx_in - unrot_w_in / 2)
    top = Inches(cy_in - unrot_h_in / 2)

    width = Inches(unrot_w_in)
    height = Inches(unrot_h_in)

    box = slide.shapes.add_textbox(
        left,
        top,
        width,
        height
    )

    tf = box.text_frame

    tf.word_wrap = False
    tf.auto_size = MSO_AUTO_SIZE.NONE
    tf.vertical_anchor = MSO_ANCHOR.MIDDLE

    tf.margin_left = 0
    tf.margin_right = 0
    tf.margin_top = 0
    tf.margin_bottom = 0

    run = tf.paragraphs[0].add_run()

    run.text = text
    run.font.name = FONT_NAME
    run.font.size = Pt(FONT_SIZE_PT)
    run.font.color.rgb = RGBColor(0, 0, 0)

    if vertical:
        box.rotation = -90

    return box


def save_presentation(prs, output_dir, base_name=None):
    """Save the PowerPoint presentation."""

    os.makedirs(output_dir, exist_ok=True)

    filename = (
        f"{base_name}.pptx"
        if base_name
        else f"figure_{uuid.uuid4().hex[:8]}.pptx"
    )

    out_path = os.path.join(output_dir, filename)

    prs.save(out_path)

    return out_path
